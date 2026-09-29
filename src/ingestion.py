"""Document loading, chunking, and Pinecone vector ingestion."""

from pathlib import Path
from typing import List
from hashlib import sha256
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_core.documents import Document
from docx import Document as DocxDocument
from src.config import get_settings
from src.vectorstore import get_vector_store, delete_file_from_namespace, user_namespace_for


# File extensions this pipeline can process.
SUPPORTED = {".pdf", ".txt", ".md", ".docx"}


def _load_docx(path: Path) -> List[Document]:
    """Load a .docx file into a single LangChain Document."""
    doc = DocxDocument(path)
    text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
    return [Document(page_content=text, metadata={"source": str(path), "title": path.name})]


def load_file(path: Path) -> List[Document]:
    """Load a single file into LangChain Documents based on its extension."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        docs = PyPDFLoader(str(path)).load()
    elif ext in {".txt", ".md"}:
        docs = TextLoader(str(path), encoding="utf-8").load()
    elif ext == ".docx":
        docs = _load_docx(path)
    else:
        raise ValueError(f"Unsupported file type: {ext}. Use PDF, TXT, MD, or DOCX.")

    for d in docs:
        d.metadata.setdefault("source", str(path))
        d.metadata.setdefault("title", path.name)
        d.metadata["document_name"] = path.name
    return docs


def _stable_chunk_id(path: Path, chunk: Document, position: int, namespace_prefix: str = "") -> str:
    """Stable IDs make repeat ingestion idempotent instead of creating duplicates."""
    material = f"{namespace_prefix}|{path.name}|{position}|{chunk.page_content}".encode("utf-8")
    digest = sha256(material).hexdigest()[:24]
    safe_stem = "".join(c if c.isalnum() or c in "-_" else "-" for c in path.stem)[:50]
    ns_tag = "".join(c if c.isalnum() else "" for c in namespace_prefix)[:10]
    return f"{ns_tag}-{safe_stem}-{position}-{digest}"


def ingest_file(path: Path, namespace: str | None = None, user_id: str = "") -> int:
    """Load, chunk, and upsert a single file into Pinecone with namespace isolation."""
    target_ns = (namespace or get_settings().pinecone_namespace).strip()
    docs = load_file(path)
    splitter = RecursiveCharacterTextSplitter(
        chunk_size=800,
        chunk_overlap=160,
        separators=["\n\n", "\n", ". ", " ", ""],
    )
    chunks = splitter.split_documents(docs)

    for i, chunk in enumerate(chunks):
        chunk.metadata["chunk_index"] = i
        chunk.metadata["knowledge_base"] = target_ns
        if user_id:
            chunk.metadata["user_id"] = user_id
            chunk.metadata["scope"] = "private"
        else:
            chunk.metadata["scope"] = "system"

    ids = [_stable_chunk_id(path, chunk, i, target_ns) for i, chunk in enumerate(chunks)]
    get_vector_store(namespace=target_ns).add_documents(chunks, ids=ids)
    return len(chunks)


def delete_file_from_vectorstore(filename: str, namespace: str | None = None) -> None:
    """Remove indexed vectors for a document from the specified Pinecone namespace."""
    target_ns = (namespace or get_settings().pinecone_namespace).strip()
    delete_file_from_namespace(filename, target_ns)


def ingest_directory(directory: Path, namespace: str | None = None) -> int:
    """Ingest all supported files in a directory. Returns total chunk count."""
    total = 0
    if not directory.exists():
        return 0
    for path in sorted(directory.iterdir()):
        if path.is_file() and path.suffix.lower() in SUPPORTED:
            total += ingest_file(path, namespace=namespace)
    return total


def namespace(user_id: str = "") -> str:
    """Return the user's isolated namespace, or global system namespace."""
    if user_id:
        return user_namespace_for(user_id)
    return get_settings().pinecone_namespace